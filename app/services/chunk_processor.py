import json
import logging
import re
import shutil
import subprocess
import tempfile
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from fastapi import UploadFile


logger = logging.getLogger(__name__)

CHUNKABLE_EXTENSIONS = {".pdf", ".doc", ".docx", ".wps", ".xls", ".xlsx"}
DEFAULT_MAX_CHUNK_CHARS = 200000
CHAPTER_HEADING_PATTERN = re.compile(r"^\s*(第\s*[一二三四五六七八九十百千万零〇两\d]+\s*章)\s*(.+?)\s*$")
TOC_PAGE_NUMBER_PATTERN = re.compile(r"(?:\s*[\.·…．。]{2,}\s*|\s+)\d+\s*$")
LEVEL_1_HEADING_PATTERN = re.compile(
    r"^\s*(?:"
    r"第\s*[一二三四五六七八九十百千万零〇两\d]+\s*章"
    r"|[一二三四五六七八九十百千万零〇两]{1,4}\s*、"
    r"|\d{1,2}(?:[\.、](?!\d)|\s+)(?=\S)"
    r")"
)


class ChunkProcessingError(Exception):
    pass


@dataclass
class TextLine:
    text: str
    page: int | None = None


@dataclass
class TextSection:
    title: str
    lines: list[TextLine]


@dataclass
class TocChapter:
    title: str
    ordinal: int | None
    page: int | None = None


class ChunkProcessor:
    def __init__(self, storage_root: Path = Path("storage/chunks")) -> None:
        self.storage_root = storage_root

    async def process_upload(
        self,
        upload: UploadFile,
        max_chunk_chars: int = DEFAULT_MAX_CHUNK_CHARS,
        include_text: bool = True,
    ) -> dict[str, Any]:
        filename = Path(upload.filename or "").name
        if not filename:
            raise ChunkProcessingError("上传文件名不能为空")

        chunk_job_id = str(uuid.uuid4())
        file_dir = self.storage_root / chunk_job_id / safe_path_name(Path(filename).stem)
        original_dir = file_dir / "original"
        original_dir.mkdir(parents=True, exist_ok=True)

        file_path = original_dir / filename
        with file_path.open("wb") as destination:
            while chunk := await upload.read(1024 * 1024):
                destination.write(chunk)

        return self.process_file_path(
            file_path=file_path,
            filename=filename,
            file_dir=file_dir,
            chunk_job_id=chunk_job_id,
            max_chunk_chars=max_chunk_chars,
            include_text=include_text,
        )

    def process_file_path(
        self,
        file_path: Path,
        filename: str | None = None,
        file_dir: Path | None = None,
        chunk_job_id: str | None = None,
        max_chunk_chars: int = DEFAULT_MAX_CHUNK_CHARS,
        include_text: bool = True,
    ) -> dict[str, Any]:
        filename = filename or file_path.name
        extension = Path(filename.lower()).suffix
        if extension not in CHUNKABLE_EXTENSIONS:
            raise ChunkProcessingError(f"不支持的文件格式：{extension or '无扩展名'}")

        if max_chunk_chars <= 0:
            raise ChunkProcessingError("max_chunk_chars 必须大于 0")

        chunk_job_id = chunk_job_id or str(uuid.uuid4())
        file_dir = file_dir or self.storage_root / chunk_job_id / safe_path_name(Path(filename).stem)
        chunk_dir = file_dir / "chunks"
        chunk_dir.mkdir(parents=True, exist_ok=True)

        with tempfile.TemporaryDirectory(prefix="bid-chunk-") as temp_dir_name:
            temp_dir = Path(temp_dir_name)
            if extension in {".xls", ".xlsx"}:
                chunks = self._chunk_excel(file_path, filename, extension, include_text)
                strategy = "sheet"
                warnings: list[str] = []
            else:
                if extension == ".doc":
                    lines, warnings = extract_doc_lines_with_fallback(file_path, temp_dir)
                    sections, section_warnings, strategy = split_document_sections(lines, filename)
                    warnings.extend(section_warnings)
                elif extension == ".wps":
                    lines = extract_wps_lines(file_path, temp_dir)
                    sections, warnings, strategy = split_document_sections(lines, filename)
                else:
                    lines = self._extract_text_lines(file_path, extension)
                    if extension == ".pdf":
                        sections, warnings = split_pdf_by_outline(file_path, lines)
                        if sections:
                            strategy = "toc_chapter"
                        else:
                            sections, warnings, strategy = split_document_sections(lines, filename)
                    else:
                        sections, warnings, strategy = split_document_sections(lines, filename)

                chunks = build_chunk_records(
                    sections=sections,
                    filename=filename,
                    max_chunk_chars=max_chunk_chars,
                    include_text=include_text,
                )

        result = {
            "chunk_job_id": chunk_job_id,
            "filename": filename,
            "extension": extension,
            "status": "completed",
            "chunk_strategy": strategy,
            "chunks": chunks,
            "warnings": warnings,
            "chunks_dir": chunk_dir.as_posix(),
            "chunks_file": (file_dir / "chunks.json").as_posix(),
        }
        persist_chunk_result(result, file_dir, chunk_dir)
        return result

    def _extract_text_lines(self, file_path: Path, extension: str) -> list[TextLine]:
        if extension == ".pdf":
            return extract_pdf_lines(file_path)
        return extract_docx_lines(file_path)

    def _chunk_excel(
        self,
        file_path: Path,
        filename: str,
        extension: str,
        include_text: bool,
    ) -> list[dict[str, Any]]:
        if extension == ".xlsx":
            sheets = extract_xlsx_sheets(file_path)
        else:
            sheets = extract_xls_sheets(file_path)

        chunks: list[dict[str, Any]] = []
        for index, (sheet_name, text) in enumerate(sheets, start=1):
            chunk = {
                "chunk_id": f"chunk_{index:03d}",
                "title": sheet_name,
                "level": 1,
                "order": index,
                "char_count": len(text),
                "part_index": 1,
                "part_total": 1,
                "source": {
                    "filename": filename,
                    "page_start": None,
                    "page_end": None,
                    "sheet_name": sheet_name,
                },
            }
            if include_text:
                chunk["text"] = text
            chunks.append(chunk)
        return chunks


def extract_pdf_lines(file_path: Path) -> list[TextLine]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ChunkProcessingError("PDF 支持需要安装 Python 包 pypdf") from exc

    reader = PdfReader(file_path)
    lines: list[TextLine] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for line in normalize_text_lines(text):
            lines.append(TextLine(text=line, page=page_index))
    return lines


def extract_docx_lines(file_path: Path) -> list[TextLine]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ChunkProcessingError("DOCX 支持需要安装 Python 包 python-docx") from exc

    document = Document(file_path)
    lines: list[TextLine] = []
    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            for line in normalize_text_lines(block.text):
                lines.append(TextLine(text=line))
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(TextLine(text="\t".join(cells)))
    return lines


def iter_docx_blocks(document: Any) -> Any:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.ns import qn

    def walk(parent: Any) -> Any:
        for child in parent.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, document)
            elif isinstance(child, CT_Tbl):
                yield Table(child, document)
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    yield from walk(content)

    yield from walk(document.element.body)


def extract_xlsx_sheets(file_path: Path) -> list[tuple[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ChunkProcessingError("XLSX 支持需要安装 Python 包 openpyxl") from exc

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        return [
            (sheet.title, sheet_to_text(sheet.iter_rows(values_only=True)))
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def extract_xls_sheets(file_path: Path) -> list[tuple[str, str]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ChunkProcessingError("XLS 支持需要安装 Python 包 xlrd") from exc

    workbook = xlrd.open_workbook(file_path.as_posix())
    sheets: list[tuple[str, str]] = []
    for sheet in workbook.sheets():
        rows = []
        for row_index in range(sheet.nrows):
            values = [format_cell_value(sheet.cell_value(row_index, col_index)) for col_index in range(sheet.ncols)]
            row_text = "\t".join(value for value in values if value)
            if row_text:
                rows.append(row_text)
        sheets.append((sheet.name, "\n".join(rows)))
    return sheets


def convert_doc_to_docx(file_path: Path, temp_dir: Path) -> Path:
    converted_path, error = try_convert_to_docx_with_soffice(file_path, temp_dir, source_extension=".doc")
    if converted_path:
        return converted_path
    raise ChunkProcessingError(error or "DOC 文件转换失败")


def extract_doc_lines_with_fallback(file_path: Path, temp_dir: Path) -> tuple[list[TextLine], list[str]]:
    converted_path, conversion_error = try_convert_doc_to_docx_with_soffice(file_path, temp_dir)
    if converted_path:
        return extract_docx_lines(converted_path), []

    try:
        lines, aspose_warnings = extract_doc_lines_with_aspose(file_path, temp_dir)
    except ChunkProcessingError as exc:
        logger.exception(
            "DOC parse error at doc_parse_failed: file=%s libreoffice_error=%s aspose_error=%s",
            file_path,
            conversion_error,
            exc,
        )
        details = []
        if conversion_error:
            details.append(f"LibreOffice：{conversion_error}")
        details.append(f"Aspose：{exc}")
        raise ChunkProcessingError(f"DOC 文件解析失败：LibreOffice 和 Aspose 均无法处理该文件；{'；'.join(details)}") from exc

    warnings = ["LibreOffice 转换失败，已使用 DOC 兜底解析"]
    warnings.extend(aspose_warnings)
    return lines, warnings


def try_convert_doc_to_docx_with_soffice(file_path: Path, temp_dir: Path) -> tuple[Path | None, str | None]:
    return try_convert_to_docx_with_soffice(file_path, temp_dir, source_extension=".doc")


def extract_wps_lines(file_path: Path, temp_dir: Path) -> list[TextLine]:
    converted_path, error = try_convert_to_docx_with_soffice(file_path, temp_dir, source_extension=".wps")
    if not converted_path:
        raise ChunkProcessingError(error or "WPS 文件转换失败")
    return extract_docx_lines(converted_path)


def try_convert_to_docx_with_soffice(
    file_path: Path,
    temp_dir: Path,
    source_extension: str,
) -> tuple[Path | None, str | None]:
    label = source_extension.lstrip(".").upper()
    soffice = resolve_soffice_path()
    if not soffice:
        return None, f"{label} 文件解析需要安装 LibreOffice，并确保 soffice 命令可用"

    output_dir = temp_dir / "converted"
    profile_dir = temp_dir / "libreoffice-profile"
    source_path = temp_dir / f"source{source_extension}"
    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(file_path, source_path)

    command = [
        soffice,
        f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--nodefault",
        "--nolockcheck",
        "--norestore",
        "--convert-to",
        "docx",
        "--outdir",
        output_dir.as_posix(),
        source_path.as_posix(),
    ]
    try:
        subprocess.run(command, check=True, capture_output=True, text=True, timeout=120)
    except subprocess.TimeoutExpired as exc:
        logger.exception(
            "%s parse error at soffice_convert: file=%s source=%s soffice=%s profile=%s output_dir=%s error=timeout",
            label,
            file_path,
            source_path,
            soffice,
            profile_dir,
            output_dir,
        )
        return None, f"{label} 文件转换超时"
    except subprocess.CalledProcessError as exc:
        logger.exception(
            "%s parse error at soffice_convert: file=%s source=%s soffice=%s profile=%s output_dir=%s returncode=%s stdout=%s stderr=%s",
            label,
            file_path,
            source_path,
            soffice,
            profile_dir,
            output_dir,
            exc.returncode,
            (exc.stdout or "")[:3000],
            (exc.stderr or "")[:3000],
        )
        detail = (exc.stderr or exc.stdout or "").strip()
        message = f"{label} 文件转换失败"
        if detail:
            message = f"{message}：{detail[:300]}"
        converted_path = find_converted_docx(output_dir, source_path.stem)
        if converted_path:
            return converted_path, None
        return None, message

    converted_path = find_converted_docx(output_dir, source_path.stem)
    if not converted_path:
        generated = ", ".join(path.name for path in output_dir.iterdir())
        return None, f"{label} 文件转换失败：未生成 docx 文件；output_dir={output_dir} generated={generated or '空'}"
    return converted_path, None


def resolve_soffice_path() -> str | None:
    candidates = [
        shutil.which("soffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/Applications/LibreOfficeDev.app/Contents/MacOS/soffice",
        (
            Path.home()
            / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
        ).as_posix(),
        (
            Path.home()
            / ".cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice"
        ).as_posix(),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    return None


def find_converted_docx(output_dir: Path, stem: str) -> Path | None:
    converted_path = output_dir / f"{stem}.docx"
    if converted_path.exists() and converted_path.stat().st_size > 0:
        return converted_path

    candidates = [path for path in output_dir.glob("*.docx") if path.stat().st_size > 0]
    return candidates[0] if candidates else None


def extract_doc_lines_with_aspose(file_path: Path, temp_dir: Path) -> tuple[list[TextLine], list[str]]:
    try:
        import aspose.words as aw
    except ImportError as exc:
        raise ChunkProcessingError("DOC 兜底解析需要安装 Python 包 aspose-words") from exc

    output_path = temp_dir / f"{safe_path_name(file_path.stem)}.txt"
    try:
        document = aw.Document(file_path.as_posix())
        document.save(output_path.as_posix(), aw.SaveFormat.TEXT)
    except Exception as exc:
        logger.exception(
            "DOC parse error at aspose_fallback: file=%s output=%s error=%s",
            file_path,
            output_path,
            exc,
        )
        raise ChunkProcessingError(f"Aspose 读取 DOC 失败：{str(exc)[:300]}") from exc

    text = output_path.read_text(encoding="utf-8", errors="ignore")
    warnings: list[str] = []
    if "This document was truncated here because it was created in the Evaluation Mode" in text:
        logger.error(
            "DOC parse error at aspose_fallback: file=%s output=%s error=%s",
            file_path,
            output_path,
            "Aspose evaluation output was truncated",
            stack_info=True,
        )
        raise ChunkProcessingError("Aspose 评估版输出被截断，需要配置 Aspose license 或更换 DOC 解析工具")

    if "Evaluation Only. Created with Aspose.Words" in text:
        warnings.append("Aspose 兜底解析输出包含评估版水印，请注意后续许可证配置")
        text = re.sub(r"Created with an evaluation copy of Aspose\\.Words\\..*", "", text)
        text = re.sub(r"Evaluation Only\\. Created with Aspose\\.Words\\..*", "", text)

    lines = [TextLine(text=line) for line in normalize_text_lines(text)]
    if not lines:
        raise ChunkProcessingError("Aspose 未提取到可用文本")
    return lines, warnings


def split_by_level_1_headings(lines: list[TextLine], filename: str) -> tuple[list[TextSection], list[str]]:
    content_lines = [line for line in lines if line.text.strip()]
    if not content_lines:
        return [TextSection(title=Path(filename).stem, lines=[])], ["未提取到正文，已生成空切片"]

    sections: list[TextSection] = []
    current_title: str | None = None
    current_lines: list[TextLine] = []
    preface_lines: list[TextLine] = []

    for line in content_lines:
        if is_level_1_heading(line.text):
            if current_title is not None:
                sections.append(TextSection(title=current_title, lines=current_lines))
            elif preface_lines:
                sections.append(TextSection(title="前言", lines=preface_lines))

            current_title = line.text.strip()
            current_lines = [line]
            preface_lines = []
        elif current_title is None:
            preface_lines.append(line)
        else:
            current_lines.append(line)

    if current_title is not None:
        sections.append(TextSection(title=current_title, lines=current_lines))
        return sections, []

    return (
        [TextSection(title=Path(filename).stem, lines=content_lines)],
        ["未识别到一级标题，已按全文生成单个切片"],
    )


def split_by_chapter_headings(lines: list[TextLine], filename: str) -> tuple[list[TextSection], list[str]]:
    content_lines = [line for line in lines if line.text.strip()]
    if not content_lines:
        return [TextSection(title=Path(filename).stem, lines=[])], ["未提取到正文，已生成空切片"]

    sections: list[TextSection] = []
    current_title: str | None = None
    current_lines: list[TextLine] = []
    preface_lines: list[TextLine] = []

    for index, line in enumerate(content_lines):
        chapter = parse_chapter_line(line.text)
        if chapter and is_embedded_chapter_list_item(content_lines, index):
            chapter = None
        if chapter:
            if current_title is not None:
                sections.append(TextSection(title=current_title, lines=current_lines))
            elif preface_lines:
                sections.append(TextSection(title="封面", lines=preface_lines))

            current_title = chapter.title
            current_lines = [line]
            preface_lines = []
        elif current_title is None:
            preface_lines.append(line)
        else:
            current_lines.append(line)

    if current_title is None:
        return [], []

    sections.append(TextSection(title=current_title, lines=current_lines))
    return sections, []


def is_embedded_chapter_list_item(lines: list[TextLine], index: int) -> bool:
    chapter = parse_chapter_line(lines[index].text)
    if not chapter or chapter.ordinal is None:
        return False

    nearby_ordinals = []
    for nearby_index in range(max(0, index - 2), min(len(lines), index + 3)):
        nearby_chapter = parse_chapter_line(lines[nearby_index].text)
        if nearby_chapter and nearby_chapter.ordinal is not None:
            nearby_ordinals.append(nearby_chapter.ordinal)

    if len(nearby_ordinals) < 3:
        return False

    unique_ordinals = sorted(set(nearby_ordinals))
    consecutive_count = 1
    longest_consecutive = 1
    for previous, current in zip(unique_ordinals, unique_ordinals[1:]):
        if current == previous + 1:
            consecutive_count += 1
            longest_consecutive = max(longest_consecutive, consecutive_count)
        else:
            consecutive_count = 1
    return longest_consecutive >= 3


def split_document_sections(lines: list[TextLine], filename: str) -> tuple[list[TextSection], list[str], str]:
    content_lines = [line for line in lines if line.text.strip()]
    sections, warnings = split_by_toc_chapters(content_lines)
    if sections:
        return sections, warnings, "toc_chapter"

    chapter_sections, chapter_warnings = split_by_chapter_headings(content_lines, filename)
    if chapter_sections:
        return (
            chapter_sections,
            ["未识别到目录章结构，已降级为章节标题切片", *chapter_warnings],
            "chapter_heading",
        )

    fallback_sections, fallback_warnings = split_by_level_1_headings(content_lines, filename)
    return (
        fallback_sections,
        ["未识别到目录章结构，已降级为标题正则切片", *fallback_warnings],
        "level_1_heading",
    )


def split_pdf_by_outline(file_path: Path, lines: list[TextLine]) -> tuple[list[TextSection], list[str]]:
    chapters = extract_pdf_outline_chapters(file_path)
    if not chapters:
        return [], []
    return split_by_chapter_pages(lines, chapters), []


def extract_pdf_outline_chapters(file_path: Path) -> list[TocChapter]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ChunkProcessingError("PDF 支持需要安装 Python 包 pypdf") from exc

    reader = PdfReader(file_path)
    chapters: list[TocChapter] = []
    for item in getattr(reader, "outline", []) or []:
        if isinstance(item, list):
            continue

        title = str(getattr(item, "title", "")).strip()
        chapter = parse_chapter_line(title)
        if not chapter:
            continue

        try:
            page = reader.get_destination_page_number(item) + 1
        except Exception:
            continue
        chapters.append(TocChapter(title=chapter.title, ordinal=chapter.ordinal, page=page))

    return dedupe_chapters_by_ordinal(chapters)


def dedupe_chapters_by_ordinal(chapters: list[TocChapter]) -> list[TocChapter]:
    result: list[TocChapter] = []
    seen_ordinals: set[int] = set()
    for chapter in chapters:
        if chapter.ordinal is None:
            continue
        if chapter.ordinal in seen_ordinals:
            continue
        seen_ordinals.add(chapter.ordinal)
        result.append(chapter)
    return sorted(result, key=lambda chapter: chapter.ordinal or 0)


def split_by_chapter_pages(lines: list[TextLine], chapters: list[TocChapter]) -> list[TextSection]:
    sections: list[TextSection] = []
    first_page = chapters[0].page
    toc_index = find_toc_index(lines)
    if toc_index is not None:
        cover_lines = lines[:toc_index]
    else:
        cover_lines = [line for line in lines if first_page is not None and line.page is not None and line.page < first_page]
    if cover_lines:
        sections.append(TextSection(title="封面", lines=cover_lines))

    for index, chapter in enumerate(chapters):
        start_page = chapter.page
        next_page = chapters[index + 1].page if index + 1 < len(chapters) else None
        chapter_lines = [
            line
            for line in lines
            if line.page is not None
            and start_page is not None
            and line.page >= start_page
            and (next_page is None or line.page < next_page)
        ]
        sections.append(TextSection(title=chapter.title, lines=chapter_lines))
    return sections


def split_by_toc_chapters(lines: list[TextLine]) -> tuple[list[TextSection], list[str]]:
    toc_index = find_toc_index(lines)
    if toc_index is None:
        return [], []

    chapters, toc_end_index = extract_toc_chapters(lines, toc_index)
    if not chapters:
        return [], []

    body_start = min(toc_end_index + 1, len(lines))
    body_lines = lines[body_start:]
    chapter_positions = find_chapter_positions(body_lines, chapters)
    if not chapter_positions:
        return [], ["已识别到目录，但未能在正文中定位章节标题"]

    sections: list[TextSection] = []
    cover_lines = lines[:toc_index]
    if cover_lines:
        sections.append(TextSection(title="封面", lines=cover_lines))

    first_chapter_position = next((position for position in chapter_positions if position is not None), None)
    if first_chapter_position is not None and first_chapter_position > 0:
        pre_chapter_lines = body_lines[:first_chapter_position]
        if pre_chapter_lines:
            sections.append(TextSection(title=derive_between_toc_title(pre_chapter_lines), lines=pre_chapter_lines))

    warnings: list[str] = []
    missing_titles = [chapter.title for chapter, position in zip(chapters, chapter_positions) if position is None]
    if missing_titles:
        warnings.append(f"部分目录章节未能在正文中定位：{'、'.join(missing_titles[:5])}")

    for index, chapter in enumerate(chapters):
        start = chapter_positions[index]
        if start is None:
            sections.append(TextSection(title=chapter.title, lines=[]))
            continue

        next_start = next(
            (
                position
                for position in chapter_positions[index + 1 :]
                if position is not None and position > start
            ),
            len(body_lines),
        )
        sections.append(TextSection(title=chapter.title, lines=body_lines[start:next_start]))

    return sections, warnings


def derive_between_toc_title(lines: list[TextLine]) -> str:
    for line in lines:
        title = line.text.strip()
        if title and len(title) <= 80:
            return title
    return "目录后说明"


def find_toc_index(lines: list[TextLine]) -> int | None:
    for index, line in enumerate(lines[:120]):
        normalized = normalize_for_match(line.text)
        if normalized == "目录" or (normalized.endswith("目录") and len(normalized) <= 8):
            return index
    return None


def extract_toc_chapters(lines: list[TextLine], toc_index: int) -> tuple[list[TocChapter], int]:
    toc_page = lines[toc_index].page
    chapters: list[TocChapter] = []
    non_chapter_count = 0
    last_ordinal: int | None = None
    toc_end_index = toc_index

    for index in range(toc_index + 1, min(len(lines), toc_index + 120)):
        line = lines[index]
        if toc_page is not None and line.page is not None and chapters and line.page > toc_page + 2:
            break

        chapter = parse_chapter_line(line.text)
        if chapter:
            if chapters and chapter.ordinal is not None and last_ordinal is not None and chapter.ordinal <= last_ordinal:
                break
            chapters.append(chapter)
            last_ordinal = chapter.ordinal or last_ordinal
            non_chapter_count = 0
            toc_end_index = index
            continue

        if chapters:
            if has_toc_page_number(line.text):
                non_chapter_count = 0
                toc_end_index = index
                continue

            non_chapter_count += 1
            if non_chapter_count >= 5:
                break

    return chapters, toc_end_index


def find_chapter_positions(body_lines: list[TextLine], chapters: list[TocChapter]) -> list[int | None]:
    positions: list[int | None] = []
    search_start = 0
    for chapter in chapters:
        found_index: int | None = None
        for index in range(search_start, len(body_lines)):
            if is_embedded_chapter_list_item(body_lines, index):
                continue
            if line_matches_toc_chapter(body_lines[index].text, chapter):
                found_index = index
                search_start = index + 1
                break
        positions.append(found_index)
    return positions


def parse_chapter_line(text: str) -> TocChapter | None:
    normalized_text = strip_toc_page_number(text)
    normalized_text = re.sub(r"^[★☆＊*\s]+", "", normalized_text)
    match = CHAPTER_HEADING_PATTERN.match(normalized_text)
    if not match:
        return None

    chapter_mark = re.sub(r"\s+", "", match.group(1))
    title = f"{chapter_mark} {match.group(2).strip()}"
    return TocChapter(title=title, ordinal=parse_chapter_ordinal(chapter_mark))


def line_matches_toc_chapter(text: str, chapter: TocChapter) -> bool:
    body_chapter = parse_chapter_line(text)
    if not body_chapter:
        return False
    if chapter.ordinal is not None and body_chapter.ordinal is not None and chapter.ordinal != body_chapter.ordinal:
        return False

    toc_normalized = normalize_for_match(chapter.title)
    body_normalized = normalize_for_match(body_chapter.title)
    return body_normalized.startswith(toc_normalized) or toc_normalized.startswith(body_normalized)


def strip_toc_page_number(text: str) -> str:
    stripped = text.strip()
    stripped = TOC_PAGE_NUMBER_PATTERN.sub("", stripped)
    return stripped.strip()


def has_toc_page_number(text: str) -> bool:
    return bool(TOC_PAGE_NUMBER_PATTERN.search(text.strip()))


def parse_chapter_ordinal(chapter_mark: str) -> int | None:
    value = chapter_mark.replace("第", "").replace("章", "").strip()
    if value.isdigit():
        return int(value)
    return chinese_number_to_int(value)


def chinese_number_to_int(value: str) -> int | None:
    digits = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
    units = {"十": 10, "百": 100, "千": 1000, "万": 10000}
    if not value:
        return None
    if all(char in digits for char in value):
        result = 0
        for char in value:
            result = result * 10 + digits[char]
        return result

    total = 0
    current = 0
    for char in value:
        if char in digits:
            current = digits[char]
        elif char in units:
            unit = units[char]
            if current == 0:
                current = 1
            total += current * unit
            current = 0
        else:
            return None
    return total + current if total or current else None


def build_chunk_records(
    sections: list[TextSection],
    filename: str,
    max_chunk_chars: int,
    include_text: bool,
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    order = 1
    for section in sections:
        parts = split_section_lines(section.lines, max_chunk_chars)
        part_total = len(parts)
        for part_index, part_lines in enumerate(parts, start=1):
            text = "\n".join(line.text for line in part_lines)
            pages = [line.page for line in part_lines if line.page is not None]
            record = {
                "chunk_id": f"chunk_{order:03d}",
                "title": section.title,
                "level": 1,
                "order": order,
                "char_count": len(text),
                "part_index": part_index,
                "part_total": part_total,
                "source": {
                    "filename": filename,
                    "page_start": min(pages) if pages else None,
                    "page_end": max(pages) if pages else None,
                    "sheet_name": None,
                },
            }
            if include_text:
                record["text"] = text
            records.append(record)
            order += 1
    return records


def split_section_lines(lines: list[TextLine], max_chunk_chars: int) -> list[list[TextLine]]:
    if not lines:
        return [[]]

    parts: list[list[TextLine]] = []
    current: list[TextLine] = []
    current_size = 0
    for line in lines:
        line_size = len(line.text) + 1
        if current and current_size + line_size > max_chunk_chars:
            parts.append(current)
            current = []
            current_size = 0
        current.append(line)
        current_size += line_size

    if current:
        parts.append(current)
    return parts


def is_level_1_heading(text: str) -> bool:
    normalized = text.strip()
    if not normalized:
        return False
    if len(normalized) > 80:
        return False
    return bool(CHAPTER_HEADING_PATTERN.match(normalized) or LEVEL_1_HEADING_PATTERN.match(normalized))


def normalize_text_lines(text: str) -> list[str]:
    return [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if line.strip()]


def normalize_for_match(text: str) -> str:
    return re.sub(r"[\s\W_]+", "", text, flags=re.UNICODE)


def sheet_to_text(rows: Any) -> str:
    row_texts = []
    for row in rows:
        values = [format_cell_value(value) for value in row]
        row_text = "\t".join(value for value in values if value)
        if row_text:
            row_texts.append(row_text)
    return "\n".join(row_texts)


def format_cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def persist_chunk_result(result: dict[str, Any], file_dir: Path, chunk_dir: Path) -> None:
    chunks_file = file_dir / "chunks.json"
    chunks_file.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    for chunk in result["chunks"]:
        title = safe_path_name(chunk.get("title") or chunk["chunk_id"])
        filename = f"{chunk['chunk_id']}_{title}.md"
        chunk_path = unique_output_path(chunk_dir / filename)
        markdown = render_chunk_markdown(chunk)
        chunk_path.write_text(markdown, encoding="utf-8")


def render_chunk_markdown(chunk: dict[str, Any]) -> str:
    source = chunk.get("source") or {}
    lines = [
        f"# {chunk.get('title') or chunk['chunk_id']}",
        "",
        f"- chunk_id: {chunk['chunk_id']}",
        f"- order: {chunk['order']}",
        f"- char_count: {chunk['char_count']}",
        f"- page_start: {source.get('page_start')}",
        f"- page_end: {source.get('page_end')}",
        f"- sheet_name: {source.get('sheet_name')}",
        "",
        "## Text",
        "",
        chunk.get("text", ""),
        "",
    ]
    return "\n".join(lines)


def safe_path_name(value: str, max_length: int = 80) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", value.strip())
    cleaned = cleaned.strip("._")
    if not cleaned:
        cleaned = "file"
    return cleaned[:max_length]


def unique_output_path(path: Path) -> Path:
    if not path.exists():
        return path

    index = 1
    while True:
        candidate = path.with_name(f"{path.stem}__{index}{path.suffix}")
        if not candidate.exists():
            return candidate
        index += 1
