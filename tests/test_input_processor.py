import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
import py7zr
from docx import Document
from openpyxl import Workbook

import app.services.chunk_processor as chunk_processor_module
from app.main import app, chunk_processor, create_batch_manifest, processor, write_batch_manifest
from app.services.chunk_processor import (
    ChunkProcessingError,
    TextLine,
    TocChapter,
    extract_doc_lines_with_aspose,
    extract_docx_lines,
    extract_wps_lines,
    split_by_chapter_pages,
    split_document_sections,
)
from app.services.input_processor import decode_zip_member_name


@pytest.fixture()
def client(tmp_path: Path) -> TestClient:
    processor.storage_root = tmp_path / "uploads"
    chunk_processor.storage_root = tmp_path / "chunks"
    return TestClient(app)


def test_upload_single_pdf_returns_bid_and_file_list(client: TestClient) -> None:
    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("招标文件.pdf", b"fake pdf", "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["bid_id"]
    assert payload["status"] == "completed"

    files_response = client.get(payload["files_url"])
    assert files_response.status_code == 200
    files = files_response.json()["files"]
    assert len(files) == 1
    assert files[0]["name"] == "招标文件.pdf"
    assert files[0]["extension"] == ".pdf"


def test_upload_zip_preserves_tree(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("docs/招标文件.docx", b"docx")
        zip_file.writestr("报价.xlsx", b"xlsx")
    archive.seek(0)

    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )

    assert response.status_code == 200
    tree = client.get(response.json()["tree_url"]).json()["root"]
    assert tree["name"] == "extracted"
    assert [child["name"] for child in tree["children"]] == ["docs", "报价.xlsx"]


def test_upload_zip_keeps_wps_file(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("附件/招标公告附件三.wps", b"wps")
    archive.seek(0)

    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )

    assert response.status_code == 200
    files = client.get(response.json()["files_url"]).json()["files"]
    assert [file["relative_path"] for file in files] == ["附件/招标公告附件三.wps"]
    assert files[0]["extension"] == ".wps"

    tree = client.get(response.json()["tree_url"]).json()["root"]
    attachment_dir = next(child for child in tree["children"] if child["name"] == "附件")
    assert [child["name"] for child in attachment_dir["children"]] == ["招标公告附件三.wps"]


def test_decodes_gbk_zip_member_name() -> None:
    info = zipfile.ZipInfo("招标文件.pdf".encode("gbk").decode("cp437"))
    info.flag_bits = 0

    assert decode_zip_member_name(info) == "招标文件.pdf"


def test_keeps_existing_unicode_zip_member_name() -> None:
    info = zipfile.ZipInfo("资料/招标文件.pdf")
    info.flag_bits = 0

    assert decode_zip_member_name(info) == "资料/招标文件.pdf"


def test_upload_zip_skips_macos_metadata(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("__MACOSX/._招标文件.pdf", b"metadata")
        zip_file.writestr(".DS_Store", b"metadata")
        zip_file.writestr("招标文件.pdf", b"pdf")
    archive.seek(0)

    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )

    assert response.status_code == 200
    files = client.get(response.json()["files_url"]).json()["files"]
    assert [file["relative_path"] for file in files] == ["招标文件.pdf"]


def test_upload_7z_skips_office_temp_files(client: TestClient, tmp_path: Path) -> None:
    archive_path = tmp_path / "bid.7z"
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    (source_dir / "202403181810341242286402020.doc").write_bytes(b"doc")
    (source_dir / "~$2403181810341242286402020.doc").write_bytes(b"lock")
    (source_dir / "广西通信信息安全监控指挥中心维修及配套机房改造.pdf").write_bytes(b"pdf")

    with py7zr.SevenZipFile(archive_path, "w") as archive:
        archive.writeall(source_dir, arcname="20240408-项目")

    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.7z", archive_path.read_bytes(), "application/x-7z-compressed")},
    )

    assert response.status_code == 200
    files = client.get(response.json()["files_url"]).json()["files"]
    relative_paths = {file["relative_path"] for file in files}
    assert "20240408-项目/202403181810341242286402020.doc" in relative_paths
    assert "20240408-项目/广西通信信息安全监控指挥中心维修及配套机房改造.pdf" in relative_paths
    assert all("~$" not in path for path in relative_paths)


def test_upload_nested_zip_extracts_nested_contents(client: TestClient) -> None:
    nested = io.BytesIO()
    with zipfile.ZipFile(nested, "w") as zip_file:
        zip_file.writestr("inner.pdf", b"pdf")

    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as zip_file:
        zip_file.writestr("packs/nested.zip", nested.getvalue())
    outer.seek(0)

    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("outer.zip", outer.getvalue(), "application/zip")},
    )

    assert response.status_code == 200
    files = client.get(response.json()["files_url"]).json()["files"]
    relative_paths = {file["relative_path"] for file in files}
    assert "packs/nested.zip" not in relative_paths
    assert "packs/nested/inner.pdf" in relative_paths

    tree = client.get(response.json()["tree_url"]).json()["root"]
    packs = next(child for child in tree["children"] if child["name"] == "packs")
    assert [child["name"] for child in packs["children"]] == ["nested"]


def test_rejects_unsupported_extension(client: TestClient) -> None:
    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("run.exe", b"binary", "application/octet-stream")},
    )

    assert response.status_code == 400
    assert "不支持" in response.json()["detail"]


def test_rejects_zip_path_traversal(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("../escape.pdf", b"bad")
    archive.seek(0)

    response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bad.zip", archive.getvalue(), "application/zip")},
    )

    assert response.status_code == 400
    assert "不安全路径" in response.json()["detail"]


def test_chunk_docx_by_level_1_headings(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    document.add_paragraph("公告正文")
    document.add_paragraph("第二章 投标人须知")
    document.add_paragraph("须知正文")
    file_path = tmp_path / "招标文件.docx"
    document.save(file_path)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={
            "file": (
                "招标文件.docx",
                file_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["chunk_strategy"] == "chapter_heading"
    assert payload["chunk_job_id"]
    assert Path(payload["chunks_file"]).exists()
    assert Path(payload["chunks_dir"]).is_dir()
    assert "未识别到目录章结构" in payload["warnings"][0]
    assert [chunk["title"] for chunk in payload["chunks"]] == ["第一章 招标公告", "第二章 投标人须知"]
    assert "公告正文" in payload["chunks"][0]["text"]
    assert "content_blocks" not in payload["chunks"][0]
    chunk_files = sorted(Path(payload["chunks_dir"]).glob("*.md"))
    assert len(chunk_files) == 2
    assert "第一章 招标公告" in Path(payload["chunks_file"]).read_text(encoding="utf-8")
    assert "公告正文" in chunk_files[0].read_text(encoding="utf-8")
    assert "<table" not in chunk_files[0].read_text(encoding="utf-8")


def test_chunk_prefers_chapter_headings_over_numbered_sections(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("2025年四川联通成都枢纽北楼核心局房老旧低压配电系统隐患整治工程项目")
    document.add_paragraph("招 标 文 件")
    document.add_paragraph("招标编号：03-08-04A-2025-D-F-C23077")
    document.add_paragraph("招标人：中国联合网络通信有限公司成都市分公司（盖单位公章）")
    document.add_paragraph("2025年09月")
    document.add_paragraph("第一章  招标公告")
    document.add_paragraph("本招标项目为2025年四川联通成都枢纽北楼核心局房老旧低压配电系统隐患整治工程项目")
    document.add_paragraph("1．项目概况与招标内容")
    document.add_paragraph("1.1 项目概况")
    document.add_paragraph("第二章  投标人须知")
    document.add_paragraph("投标人须知前附表")
    document.add_paragraph("1．总则")
    document.add_paragraph("1.1 项目概况")
    document.add_paragraph("2.1.1 招标文件一般由以下部分组成：")
    document.add_paragraph("第一章  招标公告/投标邀请书")
    document.add_paragraph("第二章  投标人须知")
    document.add_paragraph("第三章  评标办法")
    document.add_paragraph("第四章  合同条款")
    document.add_paragraph("第五章  技术标准和要求")
    file_path = tmp_path / "四川联通.docx"
    document.save(file_path)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={
            "file": (
                "四川联通.docx",
                file_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["chunk_strategy"] == "chapter_heading"
    assert [chunk["title"] for chunk in payload["chunks"]] == ["封面", "第一章 招标公告", "第二章 投标人须知"]
    assert "招 标 文 件" in payload["chunks"][0]["text"]
    assert "1．项目概况与招标内容" in payload["chunks"][1]["text"]
    assert "1．总则" in payload["chunks"][2]["text"]
    assert "第三章  评标办法" in payload["chunks"][2]["text"]


def test_chunk_docx_by_toc_chapters(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("项目招标文件")
    document.add_paragraph("采购人：测试单位")
    document.add_paragraph("目录")
    document.add_paragraph("第一章 投标邀请")
    document.add_paragraph("第二章 供应商须知")
    document.add_paragraph("投标注意事项")
    document.add_paragraph("注意事项正文")
    document.add_paragraph("第一章 投标邀请(招标公告)")
    document.add_paragraph("一、项目基本情况")
    document.add_paragraph("1.项目名称：测试项目")
    document.add_paragraph("第二章 供应商须知")
    document.add_paragraph("1.说明：这里不应该成为一级切片")
    file_path = tmp_path / "目录招标文件.docx"
    document.save(file_path)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={
            "file": (
                "目录招标文件.docx",
                file_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["chunk_strategy"] == "toc_chapter"
    assert payload["warnings"] == []
    assert [chunk["title"] for chunk in payload["chunks"]] == [
        "封面",
        "投标注意事项",
        "第一章 投标邀请",
        "第二章 供应商须知",
    ]
    assert "项目招标文件" in payload["chunks"][0]["text"]
    assert "注意事项正文" in payload["chunks"][1]["text"]
    assert "一、项目基本情况" in payload["chunks"][2]["text"]
    assert "1.项目名称：测试项目" in payload["chunks"][2]["text"]
    assert "1.说明：这里不应该成为一级切片" in payload["chunks"][3]["text"]
    assert len(list(Path(payload["chunks_dir"]).glob("*.md"))) == 4


def test_toc_chapters_handle_starred_title_and_skip_embedded_chapter_list() -> None:
    lines = [
        TextLine("项目招标文件"),
        TextLine("目    录"),
        TextLine("第一章 招标公告\t1"),
        TextLine("第二章 投标人须知\t8"),
        TextLine("第三章 评标办法（综合评估法）\t31"),
        TextLine("第四章 合同条款\t41"),
        TextLine("★第五章 货物需求一览表及技术规格\t85"),
        TextLine("第六章 投标文件格式\t109"),
        TextLine("第一章 招标公告"),
        TextLine("第一章正文"),
        TextLine("第二章 投标人须知"),
        TextLine("招标文件组成"),
        TextLine("第一章 招标公告"),
        TextLine("第二章 投标人须知"),
        TextLine("第三章 评标办法"),
        TextLine("第四章 合同条款"),
        TextLine("第五章 技术标准和要求"),
        TextLine("第六章 投标文件格式"),
        TextLine("第二章正文"),
        TextLine("第二章正文补充1"),
        TextLine("第二章正文补充2"),
        TextLine("第二章正文补充3"),
        TextLine("第三章 评标办法（综合评估法）"),
        TextLine("第三章正文"),
        TextLine("第三章正文补充1"),
        TextLine("第三章正文补充2"),
        TextLine("第三章正文补充3"),
        TextLine("第四章 合同条款"),
        TextLine("第四章正文"),
        TextLine("第四章正文补充1"),
        TextLine("第四章正文补充2"),
        TextLine("第四章正文补充3"),
        TextLine("★第五章 货物需求一览表及技术规格"),
        TextLine("第五章正文"),
        TextLine("第五章正文补充1"),
        TextLine("第五章正文补充2"),
        TextLine("第五章正文补充3"),
        TextLine("第六章 投标文件格式"),
        TextLine("第六章正文"),
    ]

    sections, warnings, strategy = split_document_sections(lines, "四川联通.doc")

    assert strategy == "toc_chapter"
    assert warnings == []
    assert [section.title for section in sections] == [
        "封面",
        "第一章 招标公告",
        "第二章 投标人须知",
        "第三章 评标办法（综合评估法）",
        "第四章 合同条款",
        "第五章 货物需求一览表及技术规格",
        "第六章 投标文件格式",
    ]
    assert "第三章 评标办法\n第四章 合同条款" in "\n".join(line.text for line in sections[2].lines)
    assert sections[3].lines[0].text == "第三章 评标办法（综合评估法）"
    assert "第三章正文" in "\n".join(line.text for line in sections[3].lines)
    assert sections[5].lines[0].text == "★第五章 货物需求一览表及技术规格"
    assert "第五章正文" in "\n".join(line.text for line in sections[5].lines)


def test_toc_chapters_keep_scanning_past_numbered_subsections_with_pages() -> None:
    lines = [
        TextLine("中国工商银行青海省分行西宁住房公积金管理中心数字公积金项目"),
        TextLine("目录"),
        TextLine("第一章  招标公告\t5"),
        TextLine("1.  招标条件\t5"),
        TextLine("2.  项目概况与招标范围\t5"),
        TextLine("3.  投标人资格要求\t5"),
        TextLine("4.  招标文件的获取\t6"),
        TextLine("5.  投标文件的递交\t7"),
        TextLine("6.其他事项\t7"),
        TextLine("7.发布公告的媒介\t7"),
        TextLine("8.  联系方式\t7"),
        TextLine("第二章  投标人须知\t8"),
        TextLine("投标人须知前附表\t8"),
        TextLine("1.  总则\t12"),
        TextLine("1.1  招标项目概况\t12"),
        TextLine("第三章  评标办法（综合评估法）\t27"),
        TextLine("评标办法前附表\t27"),
        TextLine("第四章  合同条款及格式\t33"),
        TextLine("第五章 招标要求\t66"),
        TextLine("一、软硬件集成需求\t66"),
        TextLine("二、软硬件技术要求及技术指标\t67"),
        TextLine("服务要求说明书\t90"),
        TextLine("第六章  投标文件格式\t92"),
        TextLine("一、投标函\t94"),
        TextLine("二、开标一览表（报价表）\t95"),
        TextLine("第一章  招标公告"),
        TextLine("第一章正文"),
        TextLine("第一章正文补充1"),
        TextLine("第一章正文补充2"),
        TextLine("第二章  投标人须知"),
        TextLine("第二章正文"),
        TextLine("第二章正文补充1"),
        TextLine("第二章正文补充2"),
        TextLine("第三章  评标办法（综合评估法）"),
        TextLine("第三章正文"),
        TextLine("第三章正文补充1"),
        TextLine("第三章正文补充2"),
        TextLine("第四章  合同条款及格式"),
        TextLine("第四章正文"),
        TextLine("第四章正文补充1"),
        TextLine("第四章正文补充2"),
        TextLine("第五章 招标要求"),
        TextLine("第五章正文"),
        TextLine("第五章正文补充1"),
        TextLine("第五章正文补充2"),
        TextLine("第六章  投标文件格式"),
        TextLine("第六章正文"),
    ]

    sections, warnings, strategy = split_document_sections(lines, "工商银行项目.doc")

    assert strategy == "toc_chapter"
    assert warnings == []
    assert [section.title for section in sections] == [
        "封面",
        "第一章 招标公告",
        "第二章 投标人须知",
        "第三章 评标办法（综合评估法）",
        "第四章 合同条款及格式",
        "第五章 招标要求",
        "第六章 投标文件格式",
    ]
    assert sections[1].lines[0].text == "第一章  招标公告"
    assert sections[2].lines[0].text == "第二章  投标人须知"
    assert "1.  招标条件\t5" not in "\n".join(line.text for line in sections[1].lines)
    assert "第六章正文" in "\n".join(line.text for line in sections[6].lines)


def test_chunk_docx_preserves_table_order(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    document.add_paragraph("表格前正文")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "条款号"
    table.cell(0, 1).text = "条款名称"
    table.cell(1, 0).text = "1.1.2"
    table.cell(1, 1).text = "招标人"
    document.add_paragraph("表格后正文")
    document.add_paragraph("第二章 投标人须知")
    document.add_paragraph("第二章正文")
    file_path = tmp_path / "表格顺序.docx"
    document.save(file_path)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={
            "file": (
                "表格顺序.docx",
                file_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    first_text = payload["chunks"][0]["text"]
    assert first_text.index("表格前正文") < first_text.index("条款号\t条款名称")
    assert first_text.index("1.1.2\t招标人") < first_text.index("表格后正文")
    assert "条款号\t条款名称" not in payload["chunks"][1]["text"]


def test_extract_docx_table_skips_repeated_merged_cells(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=4)
    table.cell(0, 0).text = "条款号"
    table.cell(0, 1).text = "评审因素"
    table.cell(0, 2).text = "评审标准"
    table.cell(0, 2).merge(table.cell(0, 3))
    table.cell(1, 0).text = "2.1.1"
    table.cell(1, 1).text = "形式评审标准"
    table.cell(1, 2).text = "投标人名称"
    table.cell(1, 2).merge(table.cell(1, 3))
    file_path = tmp_path / "合并单元格.docx"
    document.save(file_path)

    lines = [line.text for line in extract_docx_lines(file_path)]

    assert lines == [
        "条款号\t评审因素\t评审标准",
        "2.1.1\t形式评审标准\t投标人名称",
    ]


def test_chunk_doc_falls_back_to_aspose_when_soffice_fails(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fake_convert(file_path: Path, temp_dir: Path) -> tuple[Path | None, str | None]:
        return None, "模拟 LibreOffice 转换失败"

    def fake_aspose(file_path: Path, temp_dir: Path) -> tuple[list[TextLine], list[str]]:
        return (
            [
                TextLine("第一章 招标公告"),
                TextLine("公告正文"),
                TextLine("第二章 投标人须知"),
                TextLine("须知正文"),
            ],
            [],
        )

    monkeypatch.setattr(chunk_processor_module, "try_convert_doc_to_docx_with_soffice", fake_convert)
    monkeypatch.setattr(chunk_processor_module, "extract_doc_lines_with_aspose", fake_aspose)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={"file": ("老格式.doc", b"legacy doc", "application/msword")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert "LibreOffice 转换失败，已使用 DOC 兜底解析" in payload["warnings"]
    assert [chunk["title"] for chunk in payload["chunks"]] == ["第一章 招标公告", "第二章 投标人须知"]
    assert "公告正文" in payload["chunks"][0]["text"]


def test_chunk_wps_converts_to_docx(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    document.add_paragraph("公告正文")
    converted_path = tmp_path / "source.docx"
    document.save(converted_path)

    def fake_convert(file_path: Path, temp_dir: Path, source_extension: str) -> tuple[Path | None, str | None]:
        assert source_extension == ".wps"
        return converted_path, None

    monkeypatch.setattr(chunk_processor_module, "try_convert_to_docx_with_soffice", fake_convert)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={"file": ("招标公告附件三.wps", b"wps", "application/octet-stream")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["extension"] == ".wps"
    assert payload["chunk_strategy"] == "chapter_heading"
    assert [chunk["title"] for chunk in payload["chunks"]] == ["第一章 招标公告"]
    assert "公告正文" in payload["chunks"][0]["text"]


def test_chunk_wps_returns_clear_error_when_conversion_fails(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    def fake_convert(file_path: Path, temp_dir: Path, source_extension: str) -> tuple[Path | None, str | None]:
        assert source_extension == ".wps"
        return None, "WPS 文件转换失败：模拟错误"

    monkeypatch.setattr(chunk_processor_module, "try_convert_to_docx_with_soffice", fake_convert)

    with pytest.raises(ChunkProcessingError, match="WPS 文件转换失败"):
        extract_wps_lines(tmp_path / "bad.wps", tmp_path)


def test_aspose_truncated_output_is_rejected(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class FakeDocument:
        def __init__(self, path: str) -> None:
            self.path = path

        def save(self, output_path: str, save_format: object) -> None:
            Path(output_path).write_text(
                "第一章 招标公告\n"
                "This document was truncated here because it was created in the Evaluation Mode.",
                encoding="utf-8",
            )

    class FakeSaveFormat:
        TEXT = object()

    class FakeAsposeWords:
        Document = FakeDocument
        SaveFormat = FakeSaveFormat

    monkeypatch.setitem(__import__("sys").modules, "aspose", type("FakeAspose", (), {"words": FakeAsposeWords})())
    monkeypatch.setitem(__import__("sys").modules, "aspose.words", FakeAsposeWords)

    with pytest.raises(ChunkProcessingError, match="评估版输出被截断"):
        extract_doc_lines_with_aspose(tmp_path / "old.doc", tmp_path)


def test_split_by_chapter_pages_uses_outline_boundaries() -> None:
    lines = [
        TextLine("封面标题", page=1),
        TextLine("目录", page=2),
        TextLine("第一章 标题", page=2),
        TextLine("第一章 正文", page=3),
        TextLine("第六章 正文标题没有被抽取出来", page=35),
        TextLine("第六章 正文内容", page=36),
        TextLine("第七章 政府采购合同主要条款", page=47),
    ]
    chapters = [
        TocChapter("第一章 投标邀请", 1, 3),
        TocChapter("第六章 评审方法和标准", 6, 35),
        TocChapter("第七章 政府采购合同主要条款", 7, 47),
    ]

    sections = split_by_chapter_pages(lines, chapters)

    assert [section.title for section in sections] == [
        "封面",
        "第一章 投标邀请",
        "第六章 评审方法和标准",
        "第七章 政府采购合同主要条款",
    ]
    assert [line.text for line in sections[0].lines] == ["封面标题"]
    assert [line.page for line in sections[2].lines] == [35, 36]


def test_chunk_docx_without_heading_returns_single_chunk(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("这是一个没有一级标题的文件")
    file_path = tmp_path / "普通文件.docx"
    document.save(file_path)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={
            "file": (
                "普通文件.docx",
                file_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert len(payload["chunks"]) == 1
    assert payload["chunks"][0]["title"] == "普通文件"
    assert "未识别到目录章结构" in payload["warnings"][0]
    assert "未识别到一级标题" in payload["warnings"][1]


def test_chunk_xlsx_by_sheet(client: TestClient, tmp_path: Path) -> None:
    workbook = Workbook()
    active = workbook.active
    active.title = "报价清单"
    active.append(["项目", "价格"])
    active.append(["施工", 100])
    second = workbook.create_sheet("技术要求")
    second.append(["要求", "说明"])
    second.append(["工期", "30天"])
    file_path = tmp_path / "报价表.xlsx"
    workbook.save(file_path)

    response = client.post(
        "/api/interpretation/v1/chunk",
        files={
            "file": (
                "报价表.xlsx",
                file_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["chunk_strategy"] == "sheet"
    assert [chunk["title"] for chunk in payload["chunks"]] == ["报价清单", "技术要求"]
    assert payload["chunks"][0]["source"]["sheet_name"] == "报价清单"
    assert "施工\t100" in payload["chunks"][0]["text"]


def test_chunk_rejects_unsupported_extension(client: TestClient) -> None:
    response = client.post(
        "/api/interpretation/v1/chunk",
        files={"file": ("slides.pptx", b"pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )

    assert response.status_code == 400
    assert "不支持" in response.json()["detail"]


def test_start_bid_chunk_job_chunks_all_chunkable_files(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    document.add_paragraph("公告正文")
    docx_path = tmp_path / "招标文件.docx"
    document.save(docx_path)

    workbook = Workbook()
    workbook.active.title = "报价清单"
    workbook.active.append(["项目", "金额"])
    workbook.active.append(["施工", 100])
    xlsx_path = tmp_path / "报价表.xlsx"
    workbook.save(xlsx_path)

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("docs/招标文件.docx", docx_path.read_bytes())
        zip_file.writestr("报价表.xlsx", xlsx_path.read_bytes())
        zip_file.writestr("说明.txt", b"ignore")
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]

    start_response = client.post(f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs")

    assert start_response.status_code == 200
    started = start_response.json()
    status_response = client.get(started["status_url"])
    result_response = client.get(started["result_url"])
    assert status_response.status_code == 200
    assert result_response.status_code == 200
    result = result_response.json()
    assert result["status"] == "completed"
    assert result["total_files"] == 2
    assert result["processed_files"] == 2
    assert [file["relative_path"] for file in result["files"]] == ["docs/招标文件.docx", "报价表.xlsx"]
    assert [file["chunk_count"] for file in result["files"]] == [1, 1]
    assert "公告正文" in result["files"][0]["chunks"][0]["text"]
    assert "施工\t100" in result["files"][1]["chunks"][0]["text"]
    assert Path(result["files"][0]["chunks_file"]).exists()
    assert Path(result["files"][1]["chunks_dir"]).is_dir()


def test_start_bid_chunk_job_chunks_selected_files_only(client: TestClient, tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    document.add_paragraph("公告正文")
    docx_path = tmp_path / "招标文件.docx"
    document.save(docx_path)

    workbook = Workbook()
    workbook.active.title = "报价清单"
    workbook.active.append(["项目", "金额"])
    workbook.active.append(["施工", 100])
    xlsx_path = tmp_path / "报价表.xlsx"
    workbook.save(xlsx_path)

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("docs/招标文件.docx", docx_path.read_bytes())
        zip_file.writestr("报价表.xlsx", xlsx_path.read_bytes())
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]

    start_response = client.post(
        f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs",
        json={"relative_paths": ["报价表.xlsx"]},
    )

    assert start_response.status_code == 200
    result = client.get(start_response.json()["result_url"]).json()
    assert result["status"] == "completed"
    assert result["total_files"] == 1
    assert result["processed_files"] == 1
    assert [file["relative_path"] for file in result["files"]] == ["报价表.xlsx"]
    assert "施工\t100" in result["files"][0]["chunks"][0]["text"]


def test_start_bid_chunk_job_rejects_empty_selection(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("招标文件.pdf", b"fake pdf")
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]

    response = client.post(
        f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs",
        json={"relative_paths": []},
    )

    assert response.status_code == 400
    assert "请选择至少一个需要切片的文件" in response.json()["detail"]


def test_start_bid_chunk_job_rejects_missing_selected_file(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("招标文件.pdf", b"fake pdf")
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]

    response = client.post(
        f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs",
        json={"relative_paths": ["不存在.pdf"]},
    )

    assert response.status_code == 400
    assert "选择的文件不存在：不存在.pdf" in response.json()["detail"]


def test_start_bid_chunk_job_rejects_unsupported_selected_file(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("说明.txt", b"ignore")
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]

    response = client.post(
        f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs",
        json={"relative_paths": ["说明.txt"]},
    )

    assert response.status_code == 400
    assert "选择的文件不支持切片：说明.txt" in response.json()["detail"]


def test_bid_chunk_job_stops_on_first_failed_file(
    client: TestClient,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    document = Document()
    document.add_paragraph("第一章 招标公告")
    docx_path = tmp_path / "招标文件.docx"
    document.save(docx_path)

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("01.docx", docx_path.read_bytes())
        zip_file.writestr("02.docx", docx_path.read_bytes())
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]
    original_process_file_path = chunk_processor.process_file_path

    def fake_process_file_path(*args: object, **kwargs: object) -> dict:
        file_path = kwargs["file_path"]
        if Path(file_path).name == "02.docx":
            raise ChunkProcessingError("模拟第二个文件失败")
        return original_process_file_path(*args, **kwargs)

    monkeypatch.setattr(chunk_processor, "process_file_path", fake_process_file_path)

    start_response = client.post(f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs")

    assert start_response.status_code == 200
    result = client.get(start_response.json()["result_url"]).json()
    assert result["status"] == "failed"
    assert result["processed_files"] == 1
    assert len(result["files"]) == 1
    assert result["error"]["relative_path"] == "02.docx"
    assert "模拟第二个文件失败" in result["error"]["message"]


def test_start_bid_chunk_job_rejects_missing_bid(client: TestClient) -> None:
    response = client.post("/api/interpretation/v1/bids/missing/chunk-jobs")

    assert response.status_code == 404
    assert "bid_id 不存在" in response.json()["detail"]


def test_start_bid_chunk_job_rejects_when_no_chunkable_files(client: TestClient) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_file:
        zip_file.writestr("说明.txt", b"ignore")
    archive.seek(0)

    upload_response = client.post(
        "/api/interpretation/v1/upload",
        files={"file": ("bid.zip", archive.getvalue(), "application/zip")},
    )
    bid_id = upload_response.json()["bid_id"]

    response = client.post(f"/api/interpretation/v1/bids/{bid_id}/chunk-jobs")

    assert response.status_code == 400
    assert "没有可切片文件" in response.json()["detail"]


def test_bid_chunk_job_result_rejects_unfinished_job(client: TestClient) -> None:
    manifest = create_batch_manifest(
        bid_id="test-bid",
        batch_chunk_job_id="test-job",
        total_files=1,
    )
    write_batch_manifest(manifest)

    response = client.get("/api/interpretation/v1/bids/test-bid/chunk-jobs/test-job/result")

    assert response.status_code == 400
    assert "尚未完成" in response.json()["detail"]
